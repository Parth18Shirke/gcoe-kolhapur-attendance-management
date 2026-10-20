import io
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_attendance_docx(report_title: str, subtitle: str, headers: list, rows: list) -> io.BytesIO:
    """
    Generates an official Microsoft Word (.docx) document with formatted tables and signatures.
    """
    doc = docx.Document()
    
    # Page Setup A4 Portrait
    for s in doc.sections:
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    # Document Header
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_header.add_run("COLLEGE ATTENDANCE MANAGEMENT SYSTEM\n")
    r_sub.font.size = Pt(9)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    r_title = p_header.add_run(f"{report_title}\n")
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    r_sub2 = p_header.add_run(subtitle)
    r_sub2.font.size = Pt(10)
    r_sub2.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph()

    # Table
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(255, 255, 255)

    for row_data in rows:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = str(val)
            p = row_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8.5)

    doc.add_paragraph()
    doc.add_paragraph()

    # Signature Block
    p_sig = doc.add_paragraph()
    p_sig.add_run("Subject Faculty: __________________          Class Coordinator: __________________          H.O.D.: __________________")
    p_sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sig.runs[0].font.size = Pt(9.5)
    p_sig.runs[0].font.bold = True

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
