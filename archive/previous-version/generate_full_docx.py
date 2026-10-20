import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_full_docx():
    docx_path = r"c:\Users\DELL\Downloads\Attendance\Attendance_Management_System_Blueprint_SRS.docx"
    doc = docx.Document()
    
    # A4 Page Setup
    for s in doc.sections:
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        
    # Header Banner
    p_title = doc.add_paragraph()
    r1 = p_title.add_run("ATTENDANCE MANAGEMENT SYSTEM\n")
    r1.font.size = Pt(22)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(15, 23, 42)
    
    r2 = p_title.add_run("Project Blueprint & Preliminary Software Requirements Specification (SRS)\n")
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(2, 132, 199)
    
    r3 = p_title.add_run("Target Institutional Deployment: Higher Education College / Academic Institution\nDocument Version: 0.1 — Preliminary Review Draft | Publication Date: September 2026")
    r3.font.size = Pt(9.5)
    r3.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_paragraph("―" * 55)
    
    # Document Control Notice
    p_notice = doc.add_paragraph()
    r_n = p_notice.add_run("IMPORTANT NOTICE: This blueprint strictly distinguishes between: (1) Confirmed Requirements from teaching staff, (2) Proposed Design Decisions by developers, (3) Open Questions / TBD requiring institutional confirmation, and (4) Optional Future Extensions.")
    r_n.font.size = Pt(9)
    r_n.font.bold = True
    r_n.font.color.rgb = RGBColor(3, 105, 161)
    
    sections = [
        ("1. Executive Summary", 
         "The Attendance Management System (AMS) is a production-oriented institutional software solution engineered to modernize student attendance recording, curricular lecture tracking, regulatory 75% threshold compliance, and official academic reporting across collegiate departments. Unlike basic toy projects, this platform is conceived as an institutional-grade management platform supporting multi-department academic structures, enforcing institutional regulations, maintaining immutable audit trails, and generating standardized documentation ready for administrative and accreditation scrutiny."),
        
        ("2. Background & Problem Statement",
         "Higher education institutions consistently experience operational inefficiencies and data integrity risks from manual paper registers:\n• Labor-Intensive Record Keeping: Manual calculation of percentages consumes teaching time and introduces errors.\n• Delayed Remedial Intervention: Students falling below the 75% threshold are identified late in the term.\n• Absence of Pedagogical Auditing: Traditional roll-calls omit what syllabus topics were delivered.\n• Disjointed Administrative Reporting: Creating official Word/A4 reports requires manual re-entry."),
        
        ("3. Core Objectives",
         "1. Enable classroom attendance recording on devices available within the classroom.\n2. Maintain verifiable lecture logs with syllabus topics/points covered.\n3. Compute real-time subject-wise and semester-wise attendance percentages.\n4. Automatically dispatch <75% warning emails to students' college email addresses.\n5. Generate publication-quality A4 printable PDF and Microsoft Word (DOCX) reports.\n6. Guarantee institutional data integrity, role-based security, and audit trails."),
        
        ("4. System Scope",
         "IN SCOPE (BASE PHASE):\n• Academic Hierarchy: Departments, Semesters, Divisions, Subjects\n• Role-Based Access Control (Student, Faculty, Academic Admin)\n• Lecture session logging with mandatory syllabus topic entries\n• Atomic attendance capture linked directly to lecture instances\n• Automated 75% threshold monitoring engine\n• SMTP email dispatcher for student warning notices\n• ISO A4-optimized printable PDF and Microsoft Word (.docx) exports\n• Faculty edit window & administrative audit logging\n\nOUT OF SCOPE (BASE PHASE):\n• Biometric fingerprint/iris hardware, facial recognition, GPS geofencing, RFID tap cards, and native mobile apps (responsive web-first architecture)."),
        
        ("5. Stakeholders & User Roles",
         "• Student: View personal attendance profile, subject breakdown, topics covered, and receive warning emails (Read-Only).\n• Teaching Faculty: Create lecture logs, record syllabus topics, record attendance for assigned classes, and export A4/Word reports.\n• Academic Administrator / HOD: Master academic catalogs, department-wide compliance oversight, post-freeze edit approvals, and consolidated reports."),
    ]
    
    for h_text, body in sections:
        h = doc.add_heading(h_text, level=1)
        h.runs[0].font.color.rgb = RGBColor(15, 23, 42)
        h.runs[0].font.size = Pt(12)
        p = doc.add_paragraph(body)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        
    # FR Table
    h = doc.add_heading("6. Confirmed Functional Requirements", level=1)
    h.runs[0].font.color.rgb = RGBColor(15, 23, 42)
    h.runs[0].font.size = Pt(12)
    
    t_fr = doc.add_table(rows=1, cols=4)
    t_fr.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title_text in enumerate(["Req ID", "Requirement Title", "Priority", "Institutional Status"]):
        cell = t_fr.rows[0].cells[i]
        cell.text = title_text
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
        cell._tc.get_or_add_tcPr().append(shd)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(8.5)
                
    fr_data = [
        ("FR-01", "Attendance Recording", "High", "CONFIRMED (Capture mechanism TBD)"),
        ("FR-02", "Monthly & Academic Reports", "High", "CONFIRMED"),
        ("FR-03", "Lecture & Topic Tracking", "High", "CONFIRMED"),
        ("FR-04", "Offline Word & A4 Reporting", "High", "CONFIRMED"),
        ("FR-05", "Attendance % & <75% Email Alerts", "High", "CONFIRMED"),
    ]
    for r_id, title_text, prio, stat in fr_data:
        row = t_fr.add_row()
        row.cells[0].text = r_id
        row.cells[1].text = title_text
        row.cells[2].text = prio
        row.cells[3].text = stat
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    
    # Open Questions Table
    h = doc.add_heading("7. Open Questions & Clarifications Required (TBD)", level=1)
    h.runs[0].font.color.rgb = RGBColor(15, 23, 42)
    h.runs[0].font.size = Pt(12)
    
    t_q = doc.add_table(rows=1, cols=3)
    t_q.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title_text in enumerate(["Domain", "Clarification Question", "Options / Impact"]):
        cell = t_q.rows[0].cells[i]
        cell.text = title_text
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F172A"/>')
        cell._tc.get_or_add_tcPr().append(shd)
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(8.5)
                
    q_data = [
        ("Attendance Capture", "What exact workflow is intended by 'platform/device available to students'?", "(A) Faculty manual entry, (B) Student confirmation session, (C) Dynamic token."),
        ("Proxy Prevention", "How should proxy attendance be prevented if student devices are used?", "Short session countdown, one-time code, campus subnet check, or faculty review."),
        ("Corrections", "What is the policy for retroactive attendance corrections?", "Faculty 24h edit window vs. HOD approval required after session locking."),
        ("Infrastructure", "Does college have an existing Student Information System / ERP?", "Standardized CSV/Excel batch upload vs. REST API database sync."),
        ("Authentication", "Is institutional Single Sign-On (SSO) available?", "Google Workspace / M365 OAuth2 vs. standalone Django credentials."),
        ("Academic Rules", "How are cancelled lectures and official excused absences handled?", "Excluded from total applicable lectures; medical excused code support."),
        ("Email Policy", "What is the trigger schedule for <75% warning emails?", "Immediate on drop, weekly digest, monthly checkpoint batch, or manual HOD trigger."),
        ("Reports", "Is Word (DOCX), PDF, or both required as official standards?", "DOCX for office reformatting; PDF for official signed archives."),
        ("Hosting", "Where will the server be hosted?", "College on-premises Linux server vs. dedicated Cloud VPS."),
    ]
    for dom, q, opt in q_data:
        row = t_q.add_row()
        row.cells[0].text = dom
        row.cells[1].text = q
        row.cells[2].text = opt
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    
    # Next Sections
    more_sections = [
        ("8. Proposed Lecture & Topic Tracking Model",
         "The Lecture is the central operational entity in the system. Attendance is directly attached to a specific instructional lecture instance:\nSubject -> Lecture Session (Date, Time, Class, Faculty, Syllabus Topics Covered) -> Attendance Session -> Attendance Records.\nInstitutional Benefits: (1) Meets accreditation/curriculum audit requirements (NAAC/NBA), (2) Provides student transparency regarding missed concepts, and (3) Ensures instructional continuity for substitute faculty."),
        
        ("9. Attendance Calculation Methodology",
         "• Subject Attendance % = (Present Lectures in Subject / Total Applicable Lectures in Subject) * 100\n• Semester Attendance % = (Total Present Lectures Across All Subjects / Total Applicable Lectures Across All Subjects) * 100\n• 75% Rule Engine: Evaluates both subject and overall semester percentages. If either is < 75.0%, a deficit warning is flagged and sent to the student's college email."),
        
        ("10. Reporting Design & A4/Word Integration",
         "• ISO A4 Layout: Standard 210 x 297 mm pages, repeating headers, zero column truncation.\n• Microsoft Word (DOCX): Native OpenXML tables for easy administrative integration.\n• Signed PDF: Tamper-resistant format with institutional headers and signature blocks.\n• Standard Reports: Monthly Class Summaries, Semester Attendance Sheets, Subject Audit Logs, and <75% Deficit Summaries."),
        
        ("11. Proposed Production Technology Stack",
         "• Backend: Python 3.11+ / Django 5.x (Built-in ORM, RBAC, Admin UI, CSRF/XSS defense)\n• Database: PostgreSQL 16 (Strict ACID transactions, relational integrity, high concurrency)\n• Frontend: Django Templates + Bootstrap 5 + Vanilla JS (Rapid load, mobile responsive, zero build complexity)\n• API Layer: Django REST Framework (DRF) for modular future integrations\n• Documents: python-docx & WeasyPrint / ReportLab for native A4 PDF & DOCX generation\n• Deployment: Ubuntu Linux, Nginx, Gunicorn, Systemd, automated pg_dump backups"),
        
        ("12. Risks & Mitigation Strategies",
         "• Proxy Attendance: Mitigation via short session timers, faculty roster verification, or faculty-device entry.\n• Network Failure: Mitigation via offline caching (LocalStorage) and fallback paper rosters with batch backfill.\n• Clerical Errors: Mitigation via 24-hour faculty correction window with full audit logging.\n• Email Failures: Mitigation via persistent NotificationLog table with SMTP status codes.\n• Data Loss: Mitigation via automated daily PostgreSQL backups archived to secure storage."),
        
        ("13. Development Roadmap",
         "Phase 1: Requirements & Blueprint Sign-Off with Teaching Staff\nPhase 2: Database & Security Architecture (PostgreSQL, Django RBAC)\nPhase 3: Core Module Implementation (Master Data, Lecture Logging, Attendance Capture)\nPhase 4: Calculation, Reporting & Email Subsystems (A4 PDF/DOCX, SMTP Warnings)\nPhase 5: Testing, Verification & Departmental Pilot Run\nPhase 6: Full Institutional Deployment & Training"),
    ]
    
    for h_text, body in more_sections:
        h = doc.add_heading(h_text, level=1)
        h.runs[0].font.color.rgb = RGBColor(15, 23, 42)
        h.runs[0].font.size = Pt(12)
        p = doc.add_paragraph(body)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        
    doc.save(docx_path)
    print("Full DOCX generated successfully.")

if __name__ == "__main__":
    create_full_docx()
